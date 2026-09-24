"""Run the workbook-driven synthetic regression evaluation.

The suite is intentionally a scenario check, not a production accuracy claim.
It reads the same ``Inquiries`` sheet that is shipped as the dataset reference.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, Iterable, List

from openpyxl import load_workbook

from agent import run_agent
from setup_db import DEFAULT_WORKBOOK, build
from tools.db import get_db_path


def _split(value: Any) -> List[str]:
    if value is None:
        return []
    return [item.strip() for item in str(value).split("|") if item and item.strip()]


def _case_rows(workbook_path: Path) -> Iterable[Dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["Inquiries"]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        trace_header = "expected_trace_steps" if "expected_trace_steps" in headers else "expected_tools"
        required = {
            "inquiry_id", "message", "expected_intents", trace_header,
            "expected_decision", "expected_escalation", "expected_answerable_fields",
            "forbidden_claims",
        }
        missing = required - set(headers)
        if missing:
            raise ValueError("Inquiries sheet is missing columns: {}".format(", ".join(sorted(missing))))
        for values in rows:
            if not values or all(value is None for value in values):
                continue
            row = dict(zip(headers, values))
            row["_trace_header"] = trace_header
            yield row
    finally:
        workbook.close()


def _contains_forbidden(draft: str, phrases: List[str]) -> List[str]:
    folded = (draft or "").casefold()
    return [phrase for phrase in phrases if phrase.casefold() in folded]


def evaluate_case(case: Dict[str, Any]) -> Dict[str, Any]:
    result = run_agent(str(case["message"]))
    trace_names = {event.get("name") for event in result.tool_trace}
    expected_intents = _split(case.get("expected_intents"))
    expected_trace = _split(case.get(case["_trace_header"]))
    expected_escalation = _split(case.get("expected_escalation"))
    expected_facts = _split(case.get("expected_answerable_fields"))
    forbidden = _split(case.get("forbidden_claims"))
    forbidden_hits = _contains_forbidden(result.draft_response, forbidden)

    checks = {
        "decision": result.decision == str(case["expected_decision"]),
        "intent_coverage": set(expected_intents).issubset(set(result.inquiry.get("intents", []))),
        "trace_coverage": set(expected_trace).issubset(trace_names),
        "escalation_coverage": set(expected_escalation).issubset(set(result.escalation)),
        "answerable_facts": set(expected_facts).issubset(set(result.verified_facts)),
        "forbidden_claims_absent": not forbidden_hits,
        "review_gate": result.decision == "ANSWER" or result.human_review_required,
    }
    return {
        "case_id": case["inquiry_id"],
        "passed": all(checks.values()),
        "checks": checks,
        "expected_decision": case["expected_decision"],
        "actual_decision": result.decision,
        "actual_intents": result.inquiry.get("intents", []),
        "missing_trace": sorted(set(expected_trace) - trace_names),
        "missing_facts": sorted(set(expected_facts) - set(result.verified_facts)),
        "missing_escalation": sorted(set(expected_escalation) - set(result.escalation)),
        "forbidden_hits": forbidden_hits,
        "human_review_required": result.human_review_required,
    }


def run_evaluation(workbook_path: Path = DEFAULT_WORKBOOK) -> Dict[str, Any]:
    if not get_db_path().is_file():
        build(workbook_path=workbook_path)
    cases = list(_case_rows(workbook_path))
    results = [evaluate_case(case) for case in cases]
    dimensions = {
        name: sum(1 for item in results if item["checks"][name])
        for name in results[0]["checks"]
    } if results else {}
    return {
        "workbook": str(workbook_path),
        "case_count": len(results),
        "passed_cases": sum(1 for item in results if item["passed"]),
        "scenario_pass_rate": (
            sum(1 for item in results if item["passed"]) / len(results) if results else 0.0
        ),
        "dimension_pass_counts": dimensions,
        "cases": results,
    }


def main(argv: List[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--json", action="store_true", help="Print machine-readable JSON only")
    args = parser.parse_args(argv)
    try:
        summary = run_evaluation(args.workbook)
    except Exception as exc:
        print("Evaluation could not run: {}".format(exc), file=sys.stderr)
        return 2

    if args.json:
        print(json.dumps(summary, ensure_ascii=False, indent=2))
    else:
        print("Synthetic scenario evaluation")
        print("Cases: {}/{} passed ({:.0%})".format(
            summary["passed_cases"], summary["case_count"], summary["scenario_pass_rate"]
        ))
        for case in summary["cases"]:
            print("{}: {}".format(case["case_id"], "PASS" if case["passed"] else "FAIL"))
            if not case["passed"]:
                print("  checks={}".format(case["checks"]))
                print("  expected={} actual={}".format(
                    case["expected_decision"], case["actual_decision"]
                ))
        print("Dimension counts: {}".format(summary["dimension_pass_counts"]))
    return 0 if summary["passed_cases"] == summary["case_count"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
