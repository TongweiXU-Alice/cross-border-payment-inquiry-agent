"""Build the synthetic SQLite database atomically from the checked-in workbook."""

import argparse
import os
import sqlite3
import tempfile
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple


PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_DB = PROJECT_DIR / "data" / "payments.db"
DEFAULT_WORKBOOK = PROJECT_DIR / "data" / "Cross_Border_Payment_Agent_V1.xlsx"
DB = Path(os.getenv("PAYMENT_AGENT_DB", str(DEFAULT_DB))).expanduser().resolve()
WORKBOOK = DEFAULT_WORKBOOK

SHEET_COLUMNS = {
    "Transactions": (
        "transaction_id", "fl_id", "customer_id", "direction", "bank", "currency",
        "original_amount", "received_amount", "internal_status", "bank_reference",
        "created_at", "settled_at", "business_line", "related_transaction_id",
    ),
    "Bank_Status": (
        "bank_reference", "bank", "bank_status", "last_updated", "reason_code", "eta",
        "bank_note",
    ),
    "Refunds": (
        "transaction_id", "refund_type", "refund_status", "refund_amount", "refund_reason",
        "refund_date", "related_transaction_id",
    ),
    "Fees": (
        "transaction_id", "fee_type", "recorded_fee", "currency", "fee_source", "verified",
    ),
    "Documents": (
        "transaction_id", "document_type", "available", "mock_file_path", "source_system",
    ),
    "SOP": (
        "rule_id", "category", "condition_or_code", "meaning", "recommended_action",
        "source_type", "notes",
    ),
}


SCHEMA_SQL = """
PRAGMA foreign_keys=ON;
PRAGMA journal_mode=DELETE;

CREATE TABLE bank_status (
    bank_reference TEXT PRIMARY KEY NOT NULL CHECK(length(trim(bank_reference)) > 0),
    bank TEXT NOT NULL CHECK(length(trim(bank)) > 0),
    bank_status TEXT NOT NULL CHECK(bank_status IN ('PROCESSING','COMPLETED','RETURNED','UNDER_RFI','FAILED')),
    last_updated TEXT NOT NULL,
    reason_code TEXT,
    eta TEXT,
    bank_note TEXT
);

CREATE TABLE transactions (
    transaction_id TEXT PRIMARY KEY NOT NULL CHECK(length(trim(transaction_id)) > 0),
    fl_id TEXT NOT NULL UNIQUE CHECK(length(trim(fl_id)) > 0),
    customer_id TEXT NOT NULL CHECK(length(trim(customer_id)) > 0),
    direction TEXT NOT NULL CHECK(length(trim(direction)) > 0),
    bank TEXT NOT NULL CHECK(length(trim(bank)) > 0),
    currency TEXT NOT NULL CHECK(length(currency) = 3),
    original_amount NUMERIC NOT NULL CHECK(original_amount >= 0),
    received_amount NUMERIC CHECK(received_amount IS NULL OR received_amount >= 0),
    internal_status TEXT NOT NULL CHECK(internal_status IN ('PROCESSING','COMPLETED','REFUNDED','FAILED')),
    bank_reference TEXT NOT NULL,
    created_at TEXT NOT NULL,
    settled_at TEXT,
    business_line TEXT NOT NULL,
    related_transaction_id TEXT,
    FOREIGN KEY(bank_reference) REFERENCES bank_status(bank_reference) DEFERRABLE INITIALLY DEFERRED,
    FOREIGN KEY(related_transaction_id) REFERENCES transactions(transaction_id) DEFERRABLE INITIALLY DEFERRED
);

CREATE TABLE refunds (
    transaction_id TEXT PRIMARY KEY NOT NULL,
    refund_type TEXT NOT NULL,
    refund_status TEXT NOT NULL,
    refund_amount NUMERIC NOT NULL CHECK(refund_amount >= 0),
    refund_reason TEXT,
    refund_date TEXT NOT NULL,
    related_transaction_id TEXT,
    FOREIGN KEY(transaction_id) REFERENCES transactions(transaction_id) DEFERRABLE INITIALLY DEFERRED,
    FOREIGN KEY(related_transaction_id) REFERENCES transactions(transaction_id) DEFERRABLE INITIALLY DEFERRED
);

CREATE TABLE fees (
    transaction_id TEXT NOT NULL,
    fee_type TEXT NOT NULL,
    recorded_fee NUMERIC NOT NULL CHECK(recorded_fee >= 0),
    currency TEXT NOT NULL CHECK(length(currency) = 3),
    fee_source TEXT NOT NULL,
    verified INTEGER NOT NULL CHECK(verified IN (0,1)),
    PRIMARY KEY(transaction_id, fee_type, fee_source),
    FOREIGN KEY(transaction_id) REFERENCES transactions(transaction_id) DEFERRABLE INITIALLY DEFERRED
);

CREATE TABLE documents (
    transaction_id TEXT NOT NULL,
    document_type TEXT NOT NULL,
    available INTEGER NOT NULL CHECK(available IN (0,1)),
    mock_file_path TEXT,
    source_system TEXT NOT NULL,
    PRIMARY KEY(transaction_id, document_type),
    FOREIGN KEY(transaction_id) REFERENCES transactions(transaction_id) DEFERRABLE INITIALLY DEFERRED
);

CREATE TABLE sop (
    rule_id TEXT PRIMARY KEY NOT NULL,
    category TEXT NOT NULL,
    condition_or_code TEXT NOT NULL,
    meaning TEXT NOT NULL,
    recommended_action TEXT NOT NULL,
    source_type TEXT NOT NULL,
    notes TEXT
);

CREATE INDEX idx_transactions_lookup
    ON transactions(original_amount, currency, bank, created_at);
CREATE INDEX idx_transactions_customer_created
    ON transactions(customer_id, created_at DESC);
CREATE INDEX idx_transactions_bank_reference
    ON transactions(bank_reference);
CREATE INDEX idx_bank_status_state
    ON bank_status(bank_status, last_updated);
CREATE INDEX idx_refunds_related
    ON refunds(related_transaction_id);
PRAGMA user_version=2;
"""


def _datetime_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip() or None


def _money(value: Any) -> Optional[float]:
    if value is None:
        return None
    return float(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _boolean(value: Any) -> int:
    if isinstance(value, bool):
        return int(value)
    if value in (0, 1):
        return int(value)
    if isinstance(value, str) and value.strip().lower() in {"true", "false", "1", "0"}:
        return int(value.strip().lower() in {"true", "1"})
    raise ValueError("Expected a boolean workbook value, got {!r}".format(value))


def _load_sheet_rows(workbook: Any, sheet_name: str) -> List[Dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise ValueError("Workbook is missing required sheet: {}".format(sheet_name))
    iterator = workbook[sheet_name].iter_rows(values_only=True)
    try:
        raw_headers = next(iterator)
    except StopIteration as exc:
        raise ValueError("Workbook sheet is empty: {}".format(sheet_name)) from exc
    headers = tuple(str(value).strip() if value is not None else "" for value in raw_headers)
    expected = SHEET_COLUMNS[sheet_name]
    if headers != expected:
        raise ValueError(
            "Unexpected columns in {}. Expected {!r}, got {!r}".format(
                sheet_name, expected, headers
            )
        )
    result: List[Dict[str, Any]] = []
    for values in iterator:
        if not values or all(value is None for value in values):
            continue
        if len(values) != len(headers):
            raise ValueError("Unexpected row width in {}".format(sheet_name))
        result.append(dict(zip(headers, values)))
    return result


def _prepare_rows(workbook: Any) -> Dict[str, List[Tuple[Any, ...]]]:
    raw = {name: _load_sheet_rows(workbook, name) for name in SHEET_COLUMNS}

    bank_status = [tuple([
        row["bank_reference"], row["bank"], row["bank_status"],
        _datetime_text(row["last_updated"]), row["reason_code"],
        _datetime_text(row["eta"]), row["bank_note"],
    ]) for row in raw["Bank_Status"]]

    transactions = [tuple([
        row["transaction_id"], row["fl_id"], row["customer_id"], row["direction"],
        row["bank"], str(row["currency"]).upper(), _money(row["original_amount"]),
        _money(row["received_amount"]), row["internal_status"], row["bank_reference"],
        _datetime_text(row["created_at"]), _datetime_text(row["settled_at"]),
        row["business_line"], row["related_transaction_id"],
    ]) for row in raw["Transactions"]]

    refunds = [tuple([
        row["transaction_id"], row["refund_type"], row["refund_status"],
        _money(row["refund_amount"]), row["refund_reason"],
        _datetime_text(row["refund_date"]), row["related_transaction_id"],
    ]) for row in raw["Refunds"]]

    fees = [tuple([
        row["transaction_id"], row["fee_type"], _money(row["recorded_fee"]),
        str(row["currency"]).upper(), row["fee_source"], _boolean(row["verified"]),
    ]) for row in raw["Fees"]]

    documents = [tuple([
        row["transaction_id"], row["document_type"], _boolean(row["available"]),
        row["mock_file_path"], row["source_system"],
    ]) for row in raw["Documents"]]

    sop = [tuple(row[column] for column in SHEET_COLUMNS["SOP"]) for row in raw["SOP"]]
    return {
        "bank_status": bank_status,
        "transactions": transactions,
        "refunds": refunds,
        "fees": fees,
        "documents": documents,
        "sop": sop,
    }


def build(
    db_path: Optional[Path] = None,
    workbook_path: Optional[Path] = None,
    force: bool = False,
) -> Path:
    """Build to a temporary file, validate it, then atomically replace the target."""

    target = Path(db_path or DB).expanduser().resolve()
    source = Path(workbook_path or WORKBOOK).expanduser().resolve()
    if target.exists() and not force:
        raise FileExistsError("Database already exists; pass force=True to replace it.")
    if not source.is_file():
        raise FileNotFoundError("Workbook not found: {}".format(source))
    target.parent.mkdir(parents=True, exist_ok=True)

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to build the synthetic database.") from exc

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        prepared = _prepare_rows(workbook)
    finally:
        workbook.close()

    descriptor, temporary_name = tempfile.mkstemp(
        prefix=".{}-".format(target.stem),
        suffix=".tmp",
        dir=str(target.parent),
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    connection: Optional[sqlite3.Connection] = None
    try:
        connection = sqlite3.connect(str(temporary))
        connection.execute("PRAGMA foreign_keys=ON")
        connection.executescript(SCHEMA_SQL)
        connection.execute("PRAGMA defer_foreign_keys=ON")
        connection.executemany(
            "INSERT INTO bank_status VALUES (?,?,?,?,?,?,?)",
            prepared["bank_status"],
        )
        connection.executemany(
            "INSERT INTO transactions VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            prepared["transactions"],
        )
        connection.executemany(
            "INSERT INTO refunds VALUES (?,?,?,?,?,?,?)",
            prepared["refunds"],
        )
        connection.executemany(
            "INSERT INTO fees VALUES (?,?,?,?,?,?)",
            prepared["fees"],
        )
        connection.executemany(
            "INSERT INTO documents VALUES (?,?,?,?,?)",
            prepared["documents"],
        )
        connection.executemany(
            "INSERT INTO sop VALUES (?,?,?,?,?,?,?)",
            prepared["sop"],
        )
        connection.commit()
        if connection.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
            raise sqlite3.DatabaseError("SQLite integrity check failed")
        foreign_key_issues = connection.execute("PRAGMA foreign_key_check").fetchall()
        if foreign_key_issues:
            raise sqlite3.IntegrityError("SQLite foreign-key check failed")
        connection.close()
        connection = None
        os.replace(str(temporary), str(target))
    except Exception:
        if connection is not None:
            connection.close()
        temporary.unlink(missing_ok=True)
        raise

    print("Created {} from {}".format(target, source.name))
    return target


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, default=DB, help="SQLite output path")
    parser.add_argument("--workbook", type=Path, default=WORKBOOK, help="Source XLSX path")
    parser.add_argument("--force", action="store_true", help="Atomically replace an existing DB")
    args = parser.parse_args(argv)
    try:
        build(args.db, args.workbook, force=args.force)
    except (FileExistsError, FileNotFoundError, RuntimeError, ValueError, sqlite3.Error) as exc:
        parser.error(str(exc))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
